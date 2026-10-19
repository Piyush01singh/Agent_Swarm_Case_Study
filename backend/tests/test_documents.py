import pytest
import pymupdf
import docx
import io
from app.services.document_service import (
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_profile_facts_from_text
)

def test_pdf_extraction_success():
    # Create an in-memory PDF with text
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Rahul Sharma\nB.Tech Computer Science\nCGPA: 8.2\nSkills: Python, PyTorch, FastAPI")
    pdf_bytes = doc.tobytes()
    doc.close()

    text = extract_text_from_pdf(pdf_bytes)
    assert "Rahul Sharma" in text
    assert "CGPA: 8.2" in text

def test_pdf_extraction_empty_failure():
    # Create a blank PDF with no text layer
    doc = pymupdf.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(ValueError, match="no extractable text layer"):
        extract_text_from_pdf(pdf_bytes)

def test_docx_extraction_success():
    doc = docx.Document()
    doc.add_paragraph("Bachelor of Technology in Computer Science")
    doc.add_paragraph("CGPA: 8.2 / 10.0")
    doc.add_paragraph("Projects: Machine Learning, NLP")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    text = extract_text_from_docx(docx_bytes)
    assert "Bachelor of Technology" in text
    assert "CGPA: 8.2" in text

def test_parse_profile_facts_from_text():
    sample = """
    Candidate: Rahul Sharma
    Education: B.Tech in Computer Science
    Score: CGPA: 8.2 / 10.0
    Technical Competencies: Python, PyTorch, Deep Learning, Docker, SQL
    """
    facts = parse_profile_facts_from_text(sample)
    assert facts["cgpa"] == 8.2
    assert facts["degree"] == "B.Tech"
    assert "Python" in facts["skills"]
    assert "PyTorch" in facts["skills"]
